"""Lectura y normalización del Excel de radicados.

Ver monitor-procesos-spec.md, sección 4. El script solo LEE este archivo —
nunca lo escribe (si el usuario lo tiene abierto en Excel, escribir causaría
fallos o corrupción).
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass

import openpyxl

logger = logging.getLogger(__name__)

RADICADO_LARGO = 23


class ExcelError(Exception):
    """Excel no encontrado, ilegible, o sin las columnas requeridas.

    Debe abortar la corrida completa (sección 7 del spec).
    """


@dataclass(frozen=True)
class FilaProceso:
    fila: int  # número de fila en el Excel (1-indexed), para reportar errores
    radicado_crudo: str  # valor original tal cual venía en la celda
    radicado: str | None  # normalizado a 23 dígitos, o None si es inválido
    nombre: str | None
    valido: bool
    error: str | None = None


def _normalizar_radicado(valor) -> tuple[str, str | None]:
    """(radicado_normalizado, error). error es None si quedó válido.

    Si la celda vino como número (int/float) en vez de texto, Excel ya pudo
    haber perdido precisión antes de que este código la vea — un double solo
    representa ~15-17 dígitos significativos, y un radicado tiene 23. En ese
    caso no tiene sentido "normalizar": se reporta como error en vez de
    arriesgarse a consultar un radicado silenciosamente equivocado.
    """
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return "", (
            "el radicado está en formato numérico en el Excel, no texto — "
            "Excel pudo truncar los últimos dígitos por precisión. "
            "Cambiar el formato de la celda a texto y volver a escribirlo."
        )
    texto = str(valor) if valor is not None else ""
    solo_digitos = re.sub(r"\D", "", texto)
    if len(solo_digitos) != RADICADO_LARGO:
        return solo_digitos, (
            f"{len(solo_digitos)} dígitos en vez de {RADICADO_LARGO}"
        )
    return solo_digitos, None


def leer_radicados(
    ruta: str,
    hoja: str,
    columna_radicado: str,
    columna_nombre: str | None = None,
) -> list[FilaProceso]:
    """Lee la hoja indicada y devuelve una fila por cada radicado no vacío.

    No deduplica — eso lo decide el llamador con radicados_a_consultar(),
    porque dos filas con el mismo radicado pueden tener nombres distintos y
    ambas deben aparecer en la notificación.
    """
    try:
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    except FileNotFoundError as exc:
        raise ExcelError(f"No se encontró el archivo Excel: {ruta}") from exc
    except Exception as exc:  # openpyxl puede lanzar varios tipos según el problema
        raise ExcelError(f"No se pudo leer el Excel {ruta}: {exc}") from exc

    if hoja not in wb.sheetnames:
        raise ExcelError(
            f"La hoja {hoja!r} no existe en {ruta}. Hojas disponibles: {wb.sheetnames}"
        )
    ws = wb[hoja]

    filas_iter = ws.iter_rows(values_only=True)
    try:
        encabezados = next(filas_iter)
    except StopIteration:
        raise ExcelError(f"La hoja {hoja!r} de {ruta} está vacía")

    indices: dict[str, int] = {}
    for i, h in enumerate(encabezados):
        if h is None:
            continue
        indices[str(h).strip()] = i

    if columna_radicado not in indices:
        raise ExcelError(
            f"No se encontró la columna {columna_radicado!r} en {ruta}, hoja "
            f"{hoja!r}. Columnas encontradas: {list(indices)}"
        )
    idx_radicado = indices[columna_radicado]

    idx_nombre = indices.get(columna_nombre) if columna_nombre else None
    if columna_nombre and idx_nombre is None:
        logger.warning(
            "La columna de nombre %r no existe en la hoja %r; se sigue sin nombre.",
            columna_nombre,
            hoja,
        )

    resultado: list[FilaProceso] = []

    for num_fila, fila in enumerate(filas_iter, start=2):  # fila 1 es encabezado
        if fila is None or all(c is None for c in fila):
            continue  # fila vacía, se ignora

        crudo = fila[idx_radicado] if idx_radicado < len(fila) else None
        if crudo is None or str(crudo).strip() == "":
            continue  # sin radicado en esta fila, se ignora

        radicado, error = _normalizar_radicado(crudo)
        nombre = None
        if idx_nombre is not None and idx_nombre < len(fila) and fila[idx_nombre] is not None:
            nombre = str(fila[idx_nombre]).strip() or None

        resultado.append(
            FilaProceso(
                fila=num_fila,
                radicado_crudo=str(crudo).strip(),
                radicado=radicado if error is None else None,
                nombre=nombre,
                valido=error is None,
                error=error,
            )
        )

    return resultado


def radicados_a_consultar(filas: list[FilaProceso]) -> list[str]:
    """Radicados válidos, sin duplicados, en el orden en que aparecieron."""
    vistos: set[str] = set()
    orden: list[str] = []
    for f in filas:
        if f.valido and f.radicado not in vistos:
            vistos.add(f.radicado)
            orden.append(f.radicado)
    return orden


if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

    ruta = sys.argv[1] if len(sys.argv) > 1 else "procesos.xlsx"
    filas = leer_radicados(
        ruta,
        hoja="Hoja 1",
        columna_radicado="NUMERO DE PROCESO",
        columna_nombre="demandante",
    )

    validas = [f for f in filas if f.valido]
    invalidas = [f for f in filas if not f.valido]

    print(f"{len(filas)} fila(s) con radicado no vacío")
    print(f"  {len(validas)} válida(s)")
    print(f"  {len(invalidas)} inválida(s)")
    for f in invalidas:
        print(f"    fila {f.fila}: {f.radicado_crudo!r} -> {f.error}")

    unicos = radicados_a_consultar(filas)
    print(f"\n{len(unicos)} radicado(s) único(s) a consultar")
    for f in validas[:5]:
        print(f"  fila {f.fila}: radicado={f.radicado} nombre={f.nombre!r}")
